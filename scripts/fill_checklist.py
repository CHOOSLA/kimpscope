"""docs/checklist/results.json → 기입한 체크리스트(xlsx)와 RESULTS.md를 만든다.

결과의 원본은 results.json 하나다. 엑셀과 마크다운은 여기서 생성하므로 둘이 어긋나지 않는다.
실행: python scripts/fill_checklist.py   (openpyxl 필요)
"""
import json
from collections import Counter
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DIR = ROOT / 'docs' / 'checklist'
SRC = DIR / '2회차_셀프·상호_체크리스트_원본.xlsx'
OUT = DIR / '2회차_셀프·상호_체크리스트_셀프점검.xlsx'
ALLOWED = {'통과', '실패', '미확인', '해당없음'}


def link(e):
    """증빙 이름 → RESULTS.md 기준 상대 링크. evidence/ 안에 있으면 그쪽, 아니면 저장소 루트 기준."""
    if e == 'RESULTS.md':return '이 문서'
    if (DIR / 'evidence' / e).exists():return f'[{e}](evidence/{e})'
    if (DIR / e).exists():return f'[{e}]({e})'
    if (ROOT / e).exists():return f'[{e}](../../{e})'
    raise SystemExit(f'증빙 파일 없음: {e}')


def main():
    data = json.loads((DIR / 'results.json').read_text(encoding='utf-8'))
    items = data['items']
    assert len(items) == 27 and [x['no'] for x in items] == list(range(1, 28))
    assert all(x['result'] in ALLOWED for x in items)

    wb = openpyxl.load_workbook(SRC)
    ws = wb['체크리스트']
    ws['C3'], ws['C4'], ws['C5'], ws['C6'] = data['type'], data['author'], data['inspector'], data['checked_at']
    for x in items:
        row = 8 + x['no']
        assert ws.cell(row, 3).value == x['no'], f'행 {row} 번호 불일치'
        ws.cell(row, 7).value = x['result']
        ws.cell(row, 8).value = ', '.join(x.get('evidence', []))
        ws.cell(row, 9).value = x.get('note', '')
    wb.save(OUT)

    count = Counter(x['result'] for x in items)
    lines = [
        '# 2회차 셀프 체크리스트 결과',
        '',
        f"점검 유형 {data['type']} · 앱 작성자 {data['author']} · 점검자 {data['inspector']} · 점검 일시 {data['checked_at']}",
        '',
        f"대상: {data['target']}",
        '',
        f"집계 — 통과 {count['통과']} · 실패 {count['실패']} · 미확인 {count['미확인']} · 해당없음 {count['해당없음']} / 전체 27",
        '',
        '기입 기준은 양식 그대로입니다. 직접 실행해 기대한 결과를 확인한 것만 통과로 적었습니다.',
        '증빙은 `evidence/` 폴더에 있습니다. 화면 캡처는 `evidence/screenshots/`, 실행 결과는 `api_run.json`(운영 API)·`local_index.json`(로컬)입니다. 엑셀 기입본은 `2회차_셀프·상호_체크리스트_셀프점검.xlsx`입니다.',
        '',
        '| 번호 | 요구사항 | 확인 항목 | 결과 | 증빙 | 비고 |',
        '| --- | --- | --- | --- | --- | --- |',
    ]
    for x in items:
        row = 8 + x['no']
        ev = '<br>'.join(link(e) for e in x.get('evidence', []))
        cells = [str(x['no']), ws.cell(row, 4).value, ws.cell(row, 5).value, f"**{x['result']}**", ev, x.get('note', '')]
        lines.append('| ' + ' | '.join(str(c).replace('|', '\\|').replace('\n', ' ') for c in cells) + ' |')
    if data.get('sections'):
        for title, body in data['sections']:
            lines += ['', f'## {title}', '', body]
    (DIR / 'RESULTS.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(dict(count), OUT.name)


if __name__ == '__main__':
    main()
